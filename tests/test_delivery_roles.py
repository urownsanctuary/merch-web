"""Delivery decisions on isolated data; all people and phones are synthetic."""
from datetime import date
import os
from html import unescape
from io import BytesIO
import re
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect, text
from sqlalchemy.exc import IntegrityError

from tests import test_four_production_fixes as fixtures
from app import main, services, delivery_roles
from app.delivery_migration import migrate
from app.delivery_ui import action_input
from app.security import create_merchant_session, read_merchant_session


class DeliveryTests(unittest.TestCase):
    tearDown = fixtures.FourFixTests.tearDown

    def setUp(self):
        fixtures.FourFixTests.setUp(self)
        self.stack.enter_context(patch.dict(os.environ, {'DELIVERY_ROLE_FLOW_ENABLED': '1'}))
        self.names = {1: 'Сотрудник Альфа', 2: 'Сотрудник Бета', 3: 'Сотрудник Гамма'}
        with self.factory() as db:
            migrate(db, apply=True)
            db.execute(text("UPDATE merchants SET fio=:fio, fio_norm=:norm, phone=:phone, last4='0001' WHERE id=1"),
                       dict(fio=self.names[1],norm=services.fio_norm(self.names[1]),phone='7'+'0'*9+'1'))
            for mid in (2,3):
                db.execute(text("INSERT INTO merchants (id,fio,fio_norm,pass_hash,last4,phone,tu) VALUES (:id,:fio,:norm,:hash,:last4,:phone,'ТУ')"),
                    dict(id=mid,fio=self.names[mid],norm=services.fio_norm(self.names[mid]),hash=services.hash_last4(f'{mid:04d}'),last4=f'{mid:04d}',phone='7'+'0'*9+str(mid)))
            db.commit()
        self.stack.enter_context(patch('app.services.get_active_period',return_value={'year':2026,'month':7}))
        self.login(1)

    def login(self, mid):
        self.mid=mid
        token=create_merchant_session(services.fio_norm(self.names[mid]),mid)
        self.csrf=read_merchant_session(token)['csrf']
        self.client.cookies.set('merchant_session',token)

    def sql(self, sql, **params):
        with self.factory() as db:
            result=db.execute(text(sql),params)
            rows=result.mappings().all() if result.returns_rows else []
            db.commit()
            return rows

    def supply(self, day=2, boxes=10, pay_lt5=False):
        self.sql("INSERT INTO supplies VALUES ('P1',:day,:boxes)",day=date(2026,7,day),boxes=boxes)
        self.sql("UPDATE point_rates SET pay_lt5=:value",value=pay_lt5)

    def form(self, day=2, slot='DAY', adding=True):
        token=unescape(re.search('value="([^"]+)"',action_input(self.csrf,'P1',date(2026,7,day),slot,adding))[1])
        return dict(fio=self.names[self.mid],point_code='P1',day=day,slot=slot,csrf_token=self.csrf,
                    visit_action=token,confirm_non_working='1')

    def post(self, data, **changes):
        data={**data,**changes}
        endpoint='/toggle-inventory' if data['slot']=='FULL_INVENT' else '/toggle-day'
        return self.client.post(endpoint,data=data,follow_redirects=False)

    def question_form(self,response):
        self.assertEqual(response.status_code,200,response.text)
        return {k:unescape(v) for k,v in re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)"',response.text)}

    def add_first(self,day=2,slot='DAY'):
        self.login(1)
        r=self.post(self.form(day,slot))
        self.assertEqual(r.status_code,303,r.text)

    def double(self,choice='receiver',day=2,slot='DAY',mid=2):
        self.login(mid)
        confirm=self.post(self.form(day,slot))
        self.assertIn('другой мерчендайзер',confirm.text)
        form=self.question_form(confirm)
        response=self.post(form,delivery_choice='yes')
        if 'Как проходил ваш выход?' in response.text:
            form=self.question_form(response)
            response=self.post(form,delivery_choice=choice)
        return response,form

    def totals(self,mid):
        with self.factory() as db:
            return services.compute_point_total(db,mid,'P1',2026,7)

    def test_single_supply_and_no_cancel_write(self):
        self.supply();self.add_first()
        self.assertEqual(self.totals(1)['cnt_supply'],1)
        self.login(2)
        before=self.sql('SELECT * FROM delivery_events')
        response=self.post(self.form())
        self.assertIn('Нет — вернуться',response.text)
        # Cancel is a link, not a mutation.
        self.assertEqual(len(self.sql('SELECT * FROM visits')),1)
        self.assertEqual(self.sql('SELECT * FROM delivery_events'),before)
        self.assertEqual(self.totals(1)['total'],900)

    def test_second_no_supply_leaves_first_completely_unchanged(self):
        self.supply();self.add_first()
        before=self.sql('SELECT * FROM visits WHERE merchant_id=1')
        r,_=self.double('no_supply');self.assertEqual(r.status_code,303,r.text)
        self.assertEqual(self.sql('SELECT * FROM visits WHERE merchant_id=1'),before)
        self.assertEqual(self.sql('SELECT role FROM visit_payment_roles WHERE merchant_id=1')[0]['role'],'RECEIVER')
        self.assertEqual((self.totals(1)['total'],self.totals(2)['total']),(900,500))
        self.assertEqual(self.sql('SELECT * FROM point_adjustments'),[])

    def test_transfer_totals_audit_replay_and_coffee(self):
        self.supply();self.add_first()
        r,form=self.double();self.assertEqual(r.status_code,303,r.text)
        self.assertEqual(self.post(form,delivery_choice='receiver').status_code,303)
        self.assertEqual(self.post(form,delivery_choice='receiver').status_code,303)
        self.assertEqual((self.totals(1)['total'],self.totals(2)['total']),(500,900))
        self.assertEqual([self.totals(i)['coffee_cnt'] for i in (1,2)],[1,1])
        self.assertEqual(len(self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'")),1)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),2)
        self.assertEqual(self.sql('SELECT * FROM point_adjustments'),[])
        with self.factory() as db:
            report=services.get_admin_report_rows(db,2026,7)
            self.assertEqual({r['merchant_id']:r['point_total'] for r in report},{1:500,2:900})
            overlaps=services.get_intersections_rows(db,2026,7,include_calendar_days=True)
            self.assertTrue(overlaps)

    def test_absent_or_unpaid_delivery_skips_receiver_choice(self):
        for day,boxes in ((1,0),(2,3),(6,4)):
            with self.subTest(day=day):
                if boxes:self.supply(day,boxes)
                self.add_first(day)
                response,_=self.double(day=day)
                self.assertEqual(response.status_code,303,response.text)
        self.assertEqual(self.totals(2)['cnt_no_supply'],3)
        self.assertEqual(self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'"),[])

    def test_small_paid_delivery_uses_existing_pay_lt5_rule(self):
        self.supply(boxes=3,pay_lt5=True);self.add_first()
        r,_=self.double();self.assertEqual(r.status_code,303,r.text)
        self.assertEqual(self.totals(2)['cnt_supply'],1)
        self.assertEqual(self.totals(1)['cnt_no_supply'],1)

    def test_third_receiver_is_blocked_and_audited(self):
        self.supply();self.add_first();self.double()
        before=self.sql('SELECT * FROM visit_payment_roles')
        r,form=self.double(mid=3)
        self.assertEqual(r.status_code,409,r.text)
        self.assertIn('администратору',r.text)
        self.assertEqual(self.sql('SELECT * FROM visit_payment_roles'),before)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),2)
        self.assertEqual(len(self.sql("SELECT * FROM delivery_events WHERE event_type='EXCEPTION_RECEIVER_ASSIGNED'")),1)
        self.post(form,delivery_choice='receiver')
        self.assertEqual(len(self.sql("SELECT * FROM delivery_events WHERE event_type LIKE 'EXCEPTION_%'")),1)

    def test_submitted_first_is_unchanged_and_exception_visible(self):
        self.supply();self.add_first()
        self.sql("INSERT INTO monthly_submissions (merchant_id,month_key,status) VALUES (1,'2026-07-01','submitted')")
        before=self.totals(1)
        r,_=self.double();self.assertEqual(r.status_code,409,r.text)
        self.assertIn('подтверждена',r.text)
        self.assertEqual(self.totals(1),before)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),1)
        self.client.cookies.set('admin_auth',main.get_admin_cookie_value())
        page=self.client.get('/admin-delivery-exceptions?year=2026&month=7')
        self.assertIn('Сверка первого уже отправлена',page.text)

    def test_submitted_first_can_remain_unchanged_for_no_supply_second(self):
        self.supply();self.add_first()
        self.sql("INSERT INTO monthly_submissions (merchant_id,month_key,status) VALUES (1,'2026-07-01','submitted')")
        r,_=self.double('no_supply');self.assertEqual(r.status_code,303,r.text)
        self.assertEqual(self.totals(1)['total'],900)

    def test_atomic_rollback_on_audit_failure(self):
        self.supply();self.add_first();self.login(2)
        first=self.question_form(self.post(self.form()))
        second=self.question_form(self.post(first,delivery_choice='yes'))
        with patch('app.delivery_roles.set_role',side_effect=RuntimeError('synthetic failure')):
            r=self.post(second,delivery_choice='receiver')
        self.assertEqual(r.status_code,503)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),1)
        self.assertEqual([r['merchant_id'] for r in self.sql('SELECT * FROM visit_payment_roles')],[1])
        self.assertEqual(len(self.sql('SELECT * FROM delivery_events')),1)
        self.assertEqual(self.totals(1)['total'],900)

    def test_legacy_note_before_transfer_blocks_instead_of_double_deduction(self):
        self.supply();self.add_first()
        self.sql("INSERT INTO point_adjustments (merchant_id,point_code,month_key,note_amount,note_comment) VALUES (1,'P1','2026-07-01',-400,:note)",note=services.no_supply_adjustment_marker(date(2026,7,2)))
        r,_=self.double();self.assertEqual(r.status_code,409,r.text)
        self.assertEqual(self.totals(1)['total'],500)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),1)

    def test_legacy_note_after_role_is_rejected(self):
        self.supply();self.add_first();self.double();self.login(1)
        r=self.client.post('/save-point-note-no-supply',data=dict(fio=self.names[1],point_code='P1',supply_day=2,csrf_token=self.csrf),follow_redirects=False)
        self.assertEqual(r.status_code,409,r.text)
        self.assertEqual(self.sql('SELECT * FROM point_adjustments'),[])
        self.assertEqual(self.totals(1)['total'],500)

    def test_remove_readd_requires_new_transfer_and_retains_history(self):
        self.supply();self.add_first();self.double()
        self.assertEqual(self.post(self.form(adding=False)).status_code,303)
        self.assertEqual(self.totals(2)['coffee_cnt'],0)
        self.assertEqual(self.sql('SELECT * FROM visit_payment_roles WHERE merchant_id=2'),[])
        self.assertEqual(self.totals(1)['cnt_supply'],1)
        r,_=self.double();self.assertEqual(r.status_code,303,r.text)
        self.assertEqual(self.totals(2)['cnt_supply'],1)
        self.assertEqual(self.totals(1)['cnt_no_supply'],1)
        self.assertEqual(len(self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'")),2)

    def test_weekend_and_special_inventory_preserve_slots_and_week_limit(self):
        for day in (3,4,5,8):
            with self.subTest(day=day):
                if day==8:
                    self.stack.enter_context(patch('app.main.get_special_inventory_days',return_value=[date(2026,7,8)]))
                slot='DAY' if day in (5,8) else 'MORNING'
                self.supply(day);self.add_first(day,slot)
                r,_=self.double(day=day,slot=slot);self.assertEqual(r.status_code,303,r.text)
                page=self.client.get('/day-action-page',params=dict(fio=self.names[2],point_code='P1',day=day),follow_redirects=False)
                if day == 5:
                    self.assertEqual(page.status_code,303)
                    self.assertIn('/calendar-page',page.headers['location'])
                elif day == 4:
                    self.assertIn('Убрать утренний выход',page.text)
                    self.assertIn('Добавить вечерний полный инвент',page.text)
                else:
                    self.assertIn('Утренний выход',page.text);self.assertIn('Вечерний выход',page.text)
        evening=self.post(self.form(8,'FULL_INVENT'))
        form=self.question_form(evening)
        self.assertEqual(self.post(form,delivery_choice='yes').status_code,303)
        self.assertEqual(self.totals(2)['cnt_full_inv'],1)
        # Same ISO week, existing inventory limit must still reject another evening.
        r=self.post(self.form(10,'EVENING'))
        self.assertEqual(r.status_code,303,r.text)
        self.assertIn('visit_error=inventory_week',r.headers['location'])
        page=self.client.get('/calendar-page',params=dict(fio=self.names[2],point_code='P1'))
        self.assertIn('badge-day">В',page.text);self.assertIn('badge-inv">И',page.text)

    def test_csrf_tampering_and_stale_period_are_rejected(self):
        form=self.form()
        self.assertEqual(self.post(form,csrf_token='wrong').status_code,403)
        self.assertEqual(self.post(form,day=6).status_code,409)
        self.assertEqual(self.post(form,delivery_choice='receiver').status_code,409)
        self.assertEqual(self.sql('SELECT * FROM visits'),[])

    def test_unique_receiver_database_constraint(self):
        self.supply();self.add_first();self.double()
        with self.factory() as db,self.assertRaises(IntegrityError):
            delivery_roles.set_role(db,3,'P1',date(2026,7,2),'RECEIVER',123,'synthetic')

    def test_export_snapshots_survive_profile_rate_changes_and_are_private(self):
        self.supply();self.add_first();self.double()
        snapshot=self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'")[0]
        self.sql("UPDATE merchants SET fio='Изменённое Имя',phone=NULL,is_active=FALSE WHERE id=1")
        self.sql('UPDATE point_rates SET rate_supply=9999,rate_no_supply=8888')
        denied=self.client.get('/admin-export-delivery-changes?year=2026&month=7',follow_redirects=False)
        self.assertEqual(denied.status_code,303)
        self.client.cookies.set('admin_auth',main.get_admin_cookie_value())
        response=self.client.get('/admin-export-delivery-changes?year=2026&month=7')
        self.assertEqual(response.status_code,200,response.text)
        sheet=load_workbook(BytesIO(response.content)).active
        rows=list(sheet.values)
        self.assertNotIn('Причина изменения',rows[0]);self.assertEqual(len(rows),2)
        self.assertEqual(rows[1][2:9],(snapshot['first_fio'],snapshot['first_phone'],800,400,snapshot['second_fio'],snapshot['second_phone'],800))
        self.assertEqual(response.headers['cache-control'],'no-store')

    def test_same_name_uses_signed_merchant_id(self):
        self.sql('UPDATE merchants SET fio=:fio,fio_norm=:norm WHERE id=2',fio=self.names[1],norm=services.fio_norm(self.names[1]))
        self.names[2]=self.names[1]
        self.supply();self.add_first();self.double()
        self.assertEqual({r['merchant_id'] for r in self.sql('SELECT * FROM visits')},{1,2})
        self.assertEqual(self.totals(1)['cnt_no_supply'],1)


    def test_original_add_form_replay_cannot_remove_confirmed_visit(self):
        self.supply();self.add_first();self.login(2)
        original=self.form()
        confirm=self.question_form(self.post(original))
        selected=self.question_form(self.post(confirm,delivery_choice='yes'))
        self.assertEqual(self.post(selected,delivery_choice='receiver').status_code,303)
        self.assertEqual(self.post(original).status_code,303)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),2)
        self.assertEqual(len(self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'")),1)

    def test_first_deletion_without_fallback_requires_admin(self):
        self.supply();self.add_first();self.double('no_supply');self.login(1)
        # No recorded fallback recipient: keep finances intact for admin review.
        self.assertEqual(self.post(self.form(adding=False)).status_code,409)
        self.assertEqual(self.totals(1)['total'],900)
        self.assertEqual(self.totals(2)['total'],500)
        self.assertEqual(self.sql('SELECT role FROM visit_payment_roles WHERE merchant_id=1')[0]['role'],'RECEIVER')

    def test_overlap_export_marks_confirmation_without_losing_sheets(self):
        self.supply();self.add_first();self.double()
        self.client.cookies.set('admin_auth',main.get_admin_cookie_value())
        response=self.client.get('/admin-export-overlaps?year=2026&month=7')
        self.assertEqual(response.status_code,200,response.text)
        wb=load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames,['Пересечения','По дням','Итоги'])
        rows=list(wb['Пересечения'].values)
        self.assertEqual(rows[0][-1],'Подтверждение совместного выхода')
        self.assertEqual(rows[1][-1],'Подтверждён')
        self.assertEqual(wb['По дням'].max_row,2)

    def test_profile_reset_keeps_transfer_snapshots(self):
        from app.merchant_admin import delete_all_merchants_and_data
        self.supply();self.add_first();self.double()
        before=self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'")
        with self.factory() as db:
            delete_all_merchants_and_data(db);db.commit()
        self.assertEqual(self.sql("SELECT * FROM delivery_events WHERE event_type='TRANSFER'"),before)
        self.assertTrue(all(row['phone'] is None for row in self.sql('SELECT phone FROM merchants')))


    def test_evening_slots_never_receive_delivery_role(self):
        self.stack.enter_context(patch('app.main.get_special_inventory_days',return_value=[date(2026,7,19),date(2026,7,22)]))
        for day in (3,11,19,22):
            self.supply(day)
            self.add_first(day,'DAY' if day in (19,22) else 'MORNING')
            response,_=self.double(day=day,slot='FULL_INVENT' if day in (19,22) else 'EVENING')
            self.assertEqual(response.status_code,303,response.text)
        self.assertEqual(self.totals(2)['cnt_full_inv'],4)
        self.assertEqual(self.totals(2)['coffee_cnt'],0)
        self.assertEqual({r['merchant_id'] for r in self.sql('SELECT * FROM visit_payment_roles')},{1})
        self.assertEqual({r['slot'] for r in self.sql('SELECT slot FROM visits WHERE merchant_id=2')},{'EVENING','FULL_INVENT'})

    def test_ambiguous_legacy_first_visits_require_admin(self):
        self.supply();self.add_first()
        self.sql("INSERT INTO visits (merchant_id,point_code,visit_date,slot) VALUES (3,'P1','2026-07-02','DAY')")
        response,_=self.double()
        self.assertEqual(response.status_code,409,response.text)
        self.assertEqual(len(self.sql('SELECT * FROM visits')),2)
        self.assertEqual({r['merchant_id'] for r in self.sql('SELECT * FROM visit_payment_roles')},{1})
        self.assertEqual(self.sql("SELECT event_type FROM delivery_events WHERE event_type LIKE 'EXCEPTION_%'")[0]['event_type'],'EXCEPTION_AMBIGUOUS')


class DeliveryMigrationTests(unittest.TestCase):
    def test_dry_run_apply_reapply_preserves_legacy_rows(self):
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        engine=create_engine('sqlite+pysqlite:///:memory:')
        with Session(engine) as db:
            db.execute(text('CREATE TABLE merchants (id INTEGER PRIMARY KEY,last4 TEXT)'))
            db.execute(text('CREATE TABLE visits (id INTEGER PRIMARY KEY,merchant_id INTEGER,point_code TEXT,visit_date DATE,slot TEXT)'))
            db.execute(text("INSERT INTO merchants VALUES (1,'1234')"));db.commit()
            self.assertGreater(migrate(db)['operations'],0)
            self.assertFalse(inspect(db.connection()).has_table('delivery_events'))
            migrate(db,apply=True);db.commit()
            self.assertEqual(migrate(db,apply=True)['operations'],0)
            self.assertEqual(db.execute(text('SELECT id,last4,phone FROM merchants')).one(),(1,'1234',None))
        engine.dispose()
