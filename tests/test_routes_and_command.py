import os
import unittest
from datetime import date
from io import BytesIO
from unittest.mock import patch

os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite:///:memory:")
os.environ.setdefault("SECRET_SALT", "test-secret-salt-at-least-24-characters")
os.environ.setdefault("SESSION_SECRET", "test-session-secret-at-least-24-characters")
os.environ.setdefault("ADMIN_LOGIN", "admin")
os.environ.setdefault("ADMIN_PASSWORD", "strong-password")
os.environ.setdefault("ENVIRONMENT", "test")

from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import load_workbook

from app.db import engine
from app.legacy_migration import build_migration_plan, run_migration
from app.main import app, get_admin_cookie_value, get_db, require_draft_month
from app.security import create_merchant_session, reset_request_merchant, set_request_merchant
from app.services import fio_norm, get_merchant_by_fio, hash_last4


class RouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with engine.begin() as connection:
            connection.exec_driver_sql("DROP TABLE IF EXISTS merchants")
            connection.exec_driver_sql("""
                CREATE TABLE merchants (
                    id INTEGER PRIMARY KEY, fio TEXT, fio_norm TEXT, pass_hash TEXT,
                    telegram_id TEXT, tu TEXT, created_at DATE
                )
            """)
            connection.execute(text("""
                INSERT INTO merchants (id,fio,fio_norm,pass_hash,tu)
                VALUES (1,:fio,:norm,:password,'TU-1'), (2,:other,:other_norm,:other_password,'TU-2')
            """), {
                "fio": "Иванов Иван", "norm": fio_norm("Иванов Иван"), "password": hash_last4("1234"),
                "other": "Петров Пётр", "other_norm": fio_norm("Петров Пётр"), "other_password": hash_last4("5678"),
            })

    def test_merchant_login_success_sets_session(self):
        with TestClient(app) as client:
            response = client.post("/login-page", data={"fio": "  ИВАНОВ  ИВАН ", "last4": "1234"}, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertIn("merchant_session", response.cookies)

    def test_merchant_wrong_password(self):
        with TestClient(app) as client:
            response = client.post("/login-page", data={"fio": "Иванов Иван", "last4": "9999"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("Неверные данные", response.text)

    def test_merchant_logout_clears_signed_session(self):
        with TestClient(app) as client:
            client.cookies.set("merchant_session", create_merchant_session("иванов иван"))
            response = client.get("/merchant-logout", follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers["location"], "/login-page")
        self.assertIn("merchant_session=", response.headers["set-cookie"])

    def test_one_merchant_cannot_resolve_another(self):
        token = set_request_merchant(fio_norm("Иванов Иван"))
        try:
            from sqlalchemy.orm import Session
            with Session(engine) as db:
                self.assertIsNone(get_merchant_by_fio(db, "Петров Пётр"))
                self.assertEqual(get_merchant_by_fio(db, "Иванов Иван")["id"], 1)
        finally:
            reset_request_merchant(token)

    def test_admin_login(self):
        with TestClient(app) as client:
            response = client.post(
                "/admin-login", data={"login": "admin", "password": "strong-password"}, follow_redirects=False
            )
        self.assertEqual(response.status_code, 303)
        self.assertIn("admin_auth", response.cookies)

    def test_all_three_excel_exports_exist(self):
        paths = {route.path for route in app.routes}
        self.assertTrue({
            "/admin-export-check", "/admin-export-payroll", "/admin-export-overlaps"
        }.issubset(paths))

    def test_get_toggle_routes_are_compatibility_redirects(self):
        with TestClient(app) as client:
            response = client.get(
                "/toggle-day", params={"fio": "Иванов Иван", "point_code": "P1", "day": 1},
                follow_redirects=False,
            )
        self.assertEqual(response.status_code, 303)
        self.assertIn("/calendar-page", response.headers["location"])

    def test_post_toggle_requires_csrf(self):
        with TestClient(app) as client:
            response = client.post(
                "/toggle-day",
                data={"fio": "Иванов Иван", "point_code": "P1", "day": 1, "slot": "MORNING", "csrf_token": ""},
            )
            # FastAPI 0.136 rejects an empty required Form before the handler.
            self.assertIn(response.status_code, (403, 422))
            forged = client.post(
                "/toggle-day",
                data={"fio": "Иванов Иван", "point_code": "P1", "day": 1, "slot": "MORNING", "csrf_token": "forged"},
            )
        self.assertEqual(forged.status_code, 403)

    def test_cross_site_mutation_is_rejected_before_route(self):
        with TestClient(app) as client:
            response = client.post(
                "/save-point-note-normal",
                headers={"Origin": "https://attacker.invalid", "Sec-Fetch-Site": "cross-site"},
                data={"fio": "Иванов Иван", "point_code": "P1", "note_amount": "1", "note_comment": "x"},
            )
        self.assertEqual(response.status_code, 403)

    def test_submitted_month_rejects_direct_edit(self):
        with patch(
            "app.main.compute_overall_total",
            return_value={"submission_status": "submitted"},
        ):
            with self.assertRaises(HTTPException) as raised:
                require_draft_month(None, 1, {"year": 2026, "month": 7})
        self.assertEqual(raised.exception.status_code, 409)

    def test_original_route_inventory_is_preserved(self):
        paths = {route.path for route in app.routes}
        original = {
            "/", "/db-check", "/receipts/{file_id}/{filename}", "/active-period",
            "/debug/merchants-columns", "/login", "/login-page", "/menu-page", "/point-page",
            "/merchant-logout",
            "/calendar-page", "/point-note-page", "/save-point-note-normal",
            "/save-point-note-no-supply", "/point-reimbursement-page", "/save-point-reimbursement",
            "/save-point-adjustment", "/delete-point-note", "/delete-point-reimbursement",
            "/monthly-submit-page", "/submit-monthly-submission", "/reopen-monthly-submission",
            "/day-action-page", "/toggle-day", "/toggle-inventory", "/summary-page",
            "/admin-login", "/admin-logout", "/admin-report", "/admin-data",
            "/admin-upload-supplies", "/admin-upload-rates", "/admin-upload-merchants",
            "/admin-add-merchant", "/admin-clear-month", "/admin-clear-merchants",
            "/admin-add-special-inventory-day", "/admin-delete-special-inventory-day",
            "/admin-sync-production-calendar", "/admin-calendar-override", "/admin-calendar-reset",
            "/admin-export-check", "/admin-export-payroll", "/admin-export-overlaps",
        }
        self.assertTrue(original.issubset(paths))


class ReceiptResult:
    def __init__(self, row=None):
        self.row = row

    def mappings(self):
        return self

    def first(self):
        return self.row


class ReceiptDB:
    def __init__(self, owner_norm="иванов иван"):
        self.owner_norm = owner_norm

    def execute(self, statement, params=None):
        if "FROM receipt_files rf" in str(statement):
            return ReceiptResult({
                "original_filename": "check.pdf", "content_type": "application/pdf",
                "data": b"%PDF-test", "merchant_id": 1, "fio_norm": self.owner_norm,
            })
        return ReceiptResult()

    def commit(self):
        pass

    def close(self):
        pass


class ReceiptAccessTests(unittest.TestCase):
    def setUp(self):
        self.db = ReceiptDB()
        app.dependency_overrides[get_db] = lambda: self.db

    def tearDown(self):
        app.dependency_overrides.clear()

    def test_anonymous_receipt_forbidden(self):
        with TestClient(app) as client:
            response = client.get("/receipts/file/check.pdf")
        self.assertEqual(response.status_code, 403)

    def test_owner_can_open_receipt(self):
        with TestClient(app) as client:
            client.cookies.set("merchant_session", create_merchant_session("иванов иван"))
            response = client.get("/receipts/file/check.pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b"%PDF-test")

    def test_other_merchant_forbidden(self):
        with TestClient(app) as client:
            client.cookies.set("merchant_session", create_merchant_session("петров петр"))
            response = client.get("/receipts/file/check.pdf")
        self.assertEqual(response.status_code, 403)

    def test_admin_can_open_receipt(self):
        with TestClient(app) as client:
            client.cookies.set("admin_auth", get_admin_cookie_value())
            response = client.get("/receipts/file/check.pdf")
        self.assertEqual(response.status_code, 200)


class ExcelContentTests(unittest.TestCase):
    def _get_workbook(self, path, patch_target, rows):
        with patch(patch_target, return_value=rows):
            with TestClient(app) as client:
                client.cookies.set("admin_auth", get_admin_cookie_value())
                response = client.get(path)
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content), data_only=True)

    def test_check_export_contains_complete_financial_row(self):
        row = {
            "fio": "Тестовый Мерч", "tu": "ТУ-1", "point_code": "101",
            "cnt_supply": 2, "sum_supply": 200,
            "cnt_no_supply": 1, "sum_no_supply": 70, "cnt_total_exits": 3,
            "cnt_full_inv": 1, "sum_inventory": 90,
            "coffee_cnt": 1, "coffee_sum": 10,
            "note_amount": 20, "note_comment": "Примечание",
            "reimb_amount": 30, "reimb_comment": "Возмещение",
            "reimb_receipt": "private-receipt", "has_overlap": True,
            "status": "submitted", "comment": "Месяц", "point_total": 420,
        }
        workbook = self._get_workbook(
            "/admin-export-check?year=2026&month=7",
            "app.main.get_admin_report_rows",
            [row],
        )
        sheet = workbook["Проверка"]
        headers = [cell.value for cell in sheet[1]]
        values = [cell.value for cell in sheet[2]]
        self.assertIn("Возмещение по точке (сумма)", headers)
        self.assertIn("Чек по точке", headers)
        self.assertIn("Пересечение", headers)
        self.assertEqual(values[0], "Тестовый Мерч")
        self.assertEqual(values[-1], 420)

    def test_payroll_export_has_one_rounded_row(self):
        workbook = self._get_workbook(
            "/admin-export-payroll?year=2026&month=7",
            "app.main.get_admin_payroll_rows",
            [{
                "fio": "Тестовый Мерч", "tu": "ТУ-1",
                "clean_total": 870, "payroll_total": 1000, "status": "submitted",
            }],
        )
        sheet = workbook["Ведомость"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 3).value, 870)
        self.assertEqual(sheet.cell(2, 4).value, 1000)

    def test_overlap_export_contains_canonical_slot_pair(self):
        workbook = self._get_workbook(
            "/admin-export-overlaps?year=2026&month=7",
            "app.main.get_intersections_rows",
            [{
                "visit_date": date(2026, 7, 10), "point_code": "101",
                "fio1": "А", "tu1": "ТУ-1", "slot1": "MORNING",
                "fio2": "Б", "tu2": "ТУ-2", "slot2": "MORNING",
            }],
        )
        sheet = workbook["Пересечения"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 5).value, "MORNING")
        self.assertEqual(sheet.cell(2, 8).value, "MORNING")


class MappingResult:
    def __init__(self, rows):
        self.rows = rows

    def mappings(self):
        return self

    def all(self):
        return self.rows


class MigrationDB:
    def __init__(self, rows, existing=None):
        self.rows = rows
        self.existing = existing or []
        self.inserted = []
        self.commits = 0
        self.rollbacks = 0
        self.closed = False

    def execute(self, statement, params=None):
        sql = str(statement)
        if "SELECT * FROM point_adjustments" in sql:
            return MappingResult(self.rows)
        if "SELECT legacy_key FROM point_notes" in sql:
            return MappingResult([(key,) for key in self.existing])
        if "INSERT INTO point_" in sql or "INSERT INTO reimbursement_receipts" in sql:
            self.inserted.append(params)
        return MappingResult([])

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1

    def close(self):
        self.closed = True


class MigrationCommandTests(unittest.TestCase):
    def setUp(self):
        self.rows = [{
            "id": 1, "merchant_id": 10, "point_code": "P1", "month_key": date(2026, 7, 1),
            "note_amount": 100, "note_comment": "100 ₽ — note",
            "reimb_amount": 300, "reimb_comment": "300 ₽ — reimb",
            "reimb_receipt": "receipts/a.pdf|receipts/b.png",
        }]

    def test_dry_run_does_not_commit_or_insert(self):
        db = MigrationDB(self.rows)
        with patch("app.legacy_migration.SessionLocal", return_value=db):
            report = run_migration(apply=False)
        self.assertEqual(report["mode"], "dry-run")
        self.assertEqual(db.commits, 0)
        self.assertEqual(db.inserted, [])
        self.assertEqual(db.rollbacks, 1)

    def test_apply_migrates_and_keeps_legacy(self):
        db = MigrationDB(self.rows)
        with patch("app.legacy_migration.SessionLocal", return_value=db):
            report = run_migration(apply=True)
        self.assertEqual(report["notes_to_migrate"], 1)
        self.assertEqual(report["reimbursements_to_migrate"], 1)
        self.assertEqual(report["receipts_to_migrate"], 2)
        self.assertEqual(db.commits, 1)
        self.assertFalse(any("DELETE" in str(item) for item in db.inserted))

    def test_second_apply_is_idempotent(self):
        plan = build_migration_plan(self.rows)
        existing = {
            plan.notes[0]["legacy_key"],
            plan.reimbursements[0]["legacy_key"],
            *(receipt["legacy_key"] for receipt in plan.receipts),
        }
        db = MigrationDB(self.rows, existing)
        with patch("app.legacy_migration.SessionLocal", return_value=db):
            report = run_migration(apply=True)
        self.assertEqual(report["skipped_existing"], 4)
        self.assertEqual(report["notes_to_migrate"], 0)
        self.assertEqual(db.commits, 1)

    def test_dry_run_apply_and_reapply_on_test_database(self):
        migration_engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        factory = sessionmaker(bind=migration_engine)
        with migration_engine.begin() as connection:
            connection.exec_driver_sql("""
                CREATE TABLE point_adjustments (
                    id INTEGER PRIMARY KEY, merchant_id INTEGER, point_code TEXT, month_key DATE,
                    note_amount INTEGER, note_comment TEXT, reimb_amount INTEGER,
                    reimb_comment TEXT, reimb_receipt TEXT
                )
            """)
            connection.execute(text("""
                INSERT INTO point_adjustments VALUES
                (1,10,'P1','2026-07-01',100,'100 ₽ — note',300,'300 ₽ — reimb','receipts/a.pdf|receipts/b.png')
            """))
        with patch("app.legacy_migration.SessionLocal", side_effect=factory):
            dry = run_migration(apply=False)
            applied = run_migration(apply=True)
            repeated = run_migration(apply=True)
        self.assertEqual(dry["found_legacy_rows"], 1)
        self.assertEqual(applied["notes_to_migrate"], 1)
        self.assertEqual(repeated["notes_to_migrate"], 0)
        with migration_engine.connect() as connection:
            self.assertEqual(connection.execute(text("SELECT COUNT(*) FROM point_notes")).scalar(), 1)
            self.assertEqual(connection.execute(text("SELECT COUNT(*) FROM point_reimbursements")).scalar(), 1)
            self.assertEqual(connection.execute(text("SELECT COUNT(*) FROM reimbursement_receipts")).scalar(), 2)
            self.assertEqual(connection.execute(text("SELECT COUNT(*) FROM point_adjustments")).scalar(), 1)
        migration_engine.dispose()


if __name__ == "__main__":
    unittest.main()
