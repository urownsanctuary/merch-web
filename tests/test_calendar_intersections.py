"""Calendar exits are separate from legacy inventory/slot intersections."""
import os
import unittest
import json
import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite:///:memory:")
os.environ.setdefault("SECRET_SALT", "test-secret-salt-at-least-24-characters")
os.environ.setdefault("SESSION_SECRET", "test-session-secret-at-least-24-characters")
os.environ.setdefault("ENVIRONMENT", "test")
os.environ.setdefault("ADMIN_LOGIN", "admin")
os.environ.setdefault("ADMIN_PASSWORD", "strong-password")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import main, services
from app.intersection_audit import build_report
from app import intersection_audit


class CalendarIntersectionTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool,
                                    connect_args={"check_same_thread": False})
        with self.engine.begin() as con:
            for sql in (
                "CREATE TABLE merchants (id INTEGER PRIMARY KEY, fio TEXT, tu TEXT, historical_fio TEXT, historical_tu TEXT)",
                "CREATE TABLE visits (merchant_id INTEGER, point_code TEXT, visit_date DATE, slot TEXT)",
                "CREATE TABLE point_adjustments (merchant_id INTEGER, point_code TEXT, month_key DATE, note_amount INTEGER, note_comment TEXT, reimb_amount INTEGER, reimb_comment TEXT, reimb_receipt TEXT)",
                "CREATE TABLE supplies (point_code TEXT, supply_date DATE, boxes INTEGER)",
                "CREATE TABLE point_rates (point_code TEXT, month_key DATE, rate_supply INTEGER, rate_no_supply INTEGER, rate_inventory INTEGER, coffee_enabled BOOLEAN, coffee_rate INTEGER, pay_lt5 BOOLEAN)",
                "CREATE TABLE monthly_submissions (merchant_id INTEGER, month_key DATE, status TEXT, comment TEXT, extra_amount INTEGER, receipt_path TEXT)",
                "CREATE TABLE coffee_bonus (merchant_id INTEGER, point_code TEXT, month_key DATE, days_count INTEGER)",
            ):
                con.exec_driver_sql(sql)
            con.execute(text("INSERT INTO merchants VALUES (1,:first,'TU-1',NULL,NULL),(2,:second,'TU-2',NULL,NULL),(3,:first,'TU-3',NULL,NULL)"),
                        {"first": "Сотрудник Первый", "second": "Сотрудник Второй"})
        self.db = Session(self.engine)

    def tearDown(self):
        main.app.dependency_overrides.clear()
        self.db.close()
        self.engine.dispose()

    def visit(self, person, day, slot="DAY", point="3284", month=9):
        self.db.execute(text("INSERT INTO visits VALUES (:person,:point,:day,:slot)"),
                        {"person": person, "point": point, "day": date(2026, month, day), "slot": slot})

    def fixture_3284(self):
        for day in (1, 2, 3, 4, 7, 8, 9, 10, 11):
            for person in (1, 2):
                self.visit(person, day, "MORNING" if day in (4, 11) else "DAY")
                if day in (4, 11):
                    self.visit(person, day, "EVENING")
        self.db.commit()

    def rows(self, **kwargs):
        return services.get_intersections_rows(self.db, 2026, 9, include_calendar_days=True, **kwargs)

    def test_3284_nine_dates_and_four_preserved_slot_pairs(self):
        self.fixture_3284()
        rows = self.rows()
        calendar = [r for r in rows if r["intersection_level"] == "calendar_day"]
        self.assertEqual([r["visit_date"] for r in calendar],
                         [f"2026-09-{day:02d}" for day in (1, 2, 3, 4, 7, 8, 9, 10, 11)])
        self.assertEqual(len({r["visit_date"] for r in rows}), 9)
        self.assertEqual(len(rows), 13)
        self.assertEqual({(r["visit_date"], r["slot1"]) for r in rows if r["intersection_level"] == "slot"},
                         {(f"2026-09-{day:02d}", slot) for day in (4, 11) for slot in ("MORNING", "EVENING")})
        self.assertEqual({(r["merchant_id1"], r["merchant_id2"]) for r in rows}, {(1, 2)})
        self.assertEqual({r["fio1"] for r in rows}, {"Сотрудник Первый"})
        self.assertEqual({r["fio2"] for r in rows}, {"Сотрудник Второй"})

    def test_distinct_ids_duplicates_and_multiple_presence_slots(self):
        for person in (1, 1, 2, 2, 3):
            self.visit(person, 1)
            self.visit(person, 1, "MORNING")
        self.db.commit()
        rows = self.rows()
        for level in ("calendar_day", "slot"):
            pairs = [ (r["merchant_id1"], r["merchant_id2"]) for r in rows if r["intersection_level"] == level]
            self.assertEqual(sorted(pairs), [(1, 2), (1, 3), (2, 3)])
        # Same name, different stable IDs still represents different employees.
        self.assertTrue(any(r["fio1"] == r["fio2"] for r in rows))

    def test_same_employee_never_intersects_itself(self):
        for _ in range(3):
            self.visit(1, 1)
            self.visit(1, 1, "MORNING")
        self.db.commit()
        self.assertEqual(self.rows(), [])

    def test_only_v_counts_not_supply_or_inventory(self):
        self.visit(1, 1)
        self.visit(2, 1, "FULL_INVENT")
        self.visit(1, 2, "FULL_INVENT")
        self.visit(2, 2, "FULL_INVENT")
        self.visit(1, 3, "MORNING")
        self.visit(2, 3, "EVENING")
        self.db.execute(text("INSERT INTO supplies VALUES ('3284','2026-09-01',100)"))
        self.db.commit()
        self.assertEqual(self.rows(), [])
        self.visit(2, 1, "MORNING")
        self.db.commit()
        self.assertEqual(len(self.rows()), 1)  # mixed DAY/MORNING, not same slot

    def test_point_date_period_boundaries_and_tu_filter(self):
        self.visit(1, 1)
        self.visit(2, 1, point="OTHER")
        self.visit(2, 2)
        for month in (8, 10):
            self.visit(1, 1, month=month)
            self.visit(2, 1, month=month)
        self.db.commit()
        self.assertEqual(self.rows(), [])
        self.visit(2, 1)
        self.db.commit()
        self.assertEqual(len(self.rows(tu="TU-1")), 1)
        self.assertEqual(len(self.rows(tu="TU-2")), 1)
        self.assertEqual(self.rows(tu="TU-3"), [])

    def test_no_supply_note_preserves_legacy_suppression_but_not_exit_day(self):
        self.visit(1, 4, "MORNING")
        self.visit(2, 4, "MORNING")
        self.db.execute(text("INSERT INTO point_adjustments (merchant_id,point_code,month_key,note_comment) VALUES (1,'3284','2026-09-01',:note)"),
                        {"note": services.no_supply_adjustment_marker(date(2026, 9, 4))})
        self.db.commit()
        self.assertEqual(services.get_intersections_rows(self.db, 2026, 9), [])
        self.assertEqual([r["intersection_level"] for r in self.rows()], ["calendar_day"])

    def test_registry_draft_submitted_and_payroll_unchanged(self):
        self.visit(1, 1)
        self.visit(2, 1)
        self.db.execute(text("INSERT INTO monthly_submissions (merchant_id,month_key,status) VALUES (1,'2026-09-01','draft'),(2,'2026-09-01','submitted')"))
        self.db.commit()
        with patch.object(services, "ensure_monthly_submissions_table"), patch.object(services, "ensure_point_adjustments_table"):
            after = services.get_admin_report_rows(self.db, 2026, 9)
            with patch.object(services, "_calendar_intersection_candidates", return_value=[]):
                before = services.get_admin_report_rows(self.db, 2026, 9)
            self.assertEqual([r["has_overlap"] for r in before], [False, False])
            self.assertEqual([r["has_overlap"] for r in after], [True, True])
            self.assertEqual([{k: v for k, v in r.items() if k != "has_overlap"} for r in before],
                             [{k: v for k, v in r.items() if k != "has_overlap"} for r in after])
            self.assertEqual({r["status"] for r in after}, {"draft", "submitted"})
            self.assertTrue(services.get_admin_report_rows(self.db, 2026, 9, status="submitted")[0]["has_overlap"])

    def test_excel_route_includes_days_ids_summary_and_legacy_columns(self):
        self.fixture_3284()
        main.app.dependency_overrides[main.get_db] = lambda: self.db
        with TestClient(main.app) as client:
            client.cookies.set("admin_auth", main.get_admin_cookie_value())
            response = client.get("/admin-export-overlaps?year=2026&month=9")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook["По дням"].max_row, 10)
        self.assertEqual(workbook["Пересечения"].max_row, 14)
        self.assertEqual([c.value for c in workbook["Пересечения"][1]][:8],
                         ["Дата", "Точка", "Мерч 1", "ТУ 1", "Слот 1", "Мерч 2", "ТУ 2", "Слот 2"])
        summary = dict(workbook["Итоги"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["Уникальные дни по ТТ (ТТ + дата)"], 9)
        self.assertEqual(summary["Пересечения по слотам (пары сотрудников)"], 4)

    def test_audit_is_read_only_repeatable_and_reports_duplicates(self):
        self.fixture_3284()
        self.visit(1, 1)
        self.db.commit()
        statements = []
        @event.listens_for(self.engine, "before_cursor_execute")
        def capture(connection, cursor, statement, parameters, context, executemany):
            statements.append(statement.lstrip().upper())
        report = build_report(self.db, 2026, 9)
        self.assertEqual(build_report(self.db, 2026, 9), report)
        self.assertTrue(all(s.startswith(("SELECT", "PRAGMA")) for s in statements))
        self.assertEqual(report["before"], {"points": 1, "point_dates": 2, "calendar_dates": 2})
        self.assertEqual(report["after"], {"points": 1, "point_dates": 9, "calendar_dates": 9})
        self.assertEqual(len(report["previously_missed_pairs"]), 7)
        self.assertEqual(len(report["preserved_slot_intersections"]), 4)
        self.assertEqual(report["potential_duplicate_rows"], 1)
        self.assertEqual(report["self_intersections"], 0)
        self.assertEqual(len(report["control_3284"]["dates"]), 9)


class AuditCommandSafetyTests(unittest.TestCase):
    def test_cli_enforces_read_only_before_report_and_rolls_back(self):
        fake_engine = MagicMock()
        fake_engine.dialect.name = "postgresql"
        connection = fake_engine.connect.return_value.__enter__.return_value
        transaction = connection.begin.return_value
        def report(*args):
            self.assertEqual(str(connection.execute.call_args.args[0]), "SET TRANSACTION READ ONLY")
            self.assertFalse(transaction.commit.called)
            return {"mode": "read-only"}
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "audit.json"
            argv = ["audit", "--year", "2026", "--month", "9", "--output", str(output)]
            with patch("sys.argv", argv), patch.object(intersection_audit, "create_engine", return_value=fake_engine) as factory, \
                    patch.object(intersection_audit, "Session"), patch.object(intersection_audit, "build_report", side_effect=report):
                intersection_audit.main()
            self.assertEqual(json.loads(output.read_text(encoding="utf-8")), {"mode": "read-only"})
            self.assertEqual(factory.call_args.kwargs["isolation_level"], "REPEATABLE READ")
            transaction.rollback.assert_called_once()
            transaction.commit.assert_not_called()
            connection.commit.assert_not_called()

    def test_cli_report_failure_rolls_back_and_does_not_leak_connection_details(self):
        fake_engine = MagicMock()
        fake_engine.dialect.name = "postgresql"
        connection = fake_engine.connect.return_value.__enter__.return_value
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "audit.json"
            with patch("sys.argv", ["audit", "--year", "2026", "--month", "9", "--output", str(output)]), \
                    patch.object(intersection_audit, "create_engine", return_value=fake_engine), \
                    patch.object(intersection_audit, "Session"), \
                    patch.object(intersection_audit, "build_report", side_effect=RuntimeError("sensitive-driver-details")), \
                    patch("sys.stderr") as stderr, self.assertRaises(SystemExit):
                intersection_audit.main()
            self.assertFalse(output.exists())
            self.assertNotIn("sensitive-driver-details", str(stderr.write.call_args_list))
            connection.begin.return_value.rollback.assert_called_once()
            connection.commit.assert_not_called()
