"""Query budgets and unchanged workbook output, not hardware timing thresholds."""
import asyncio
import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import event, text

from tests import test_calendar_intersections as fixtures
from app import main, services


class ExportPerformanceTests(unittest.TestCase):
    setUp = fixtures.CalendarIntersectionTests.setUp
    tearDown = fixtures.CalendarIntersectionTests.tearDown
    visit = fixtures.CalendarIntersectionTests.visit
    rows = fixtures.CalendarIntersectionTests.rows
    fixture_3284 = fixtures.CalendarIntersectionTests.fixture_3284

    def test_notes_query_count_does_not_grow_with_points(self):
        for point in range(40):
            for person in (1, 2):
                self.visit(person, 4, 'MORNING', point=str(point))
        self.db.commit()
        queries = []
        def capture(c, cursor, sql, params, ctx, many):
            queries.append(sql)
        event.listen(self.engine, 'before_cursor_execute', capture)
        try:
            with patch.object(services, 'normalized_adjustments_available', return_value=False) as schema:
                rows = self.rows()
                schema.assert_called_once()
        finally:
            event.remove(self.engine, 'before_cursor_execute', capture)
        self.assertEqual(len(rows), 40)
        self.assertEqual(len(queries), 3)  # slots, bulk notes, calendar fallback
        self.assertEqual(sum('FROM point_adjustments' in q for q in queries), 1)

    def test_normalized_bulk_notes_preserve_date_point_person_exclusions(self):
        self.db.execute(text('CREATE TABLE point_notes (merchant_id INTEGER, point_code TEXT, month_key DATE, comment TEXT)'))
        for point in ('P', 'OTHER'):
            for day in (4, 11):
                for person in (1, 2):
                    self.visit(person, day, 'MORNING', point=point)
        self.db.execute(text("INSERT INTO point_notes VALUES (1,'P','2026-09-01','unrelated'),(1,'P','2026-09-01',:note),(3,'OTHER','2026-09-01',:note),(1,'OTHER','2026-08-01',:note)"),
                        {'note': services.no_supply_adjustment_marker(services.date(2026,9,4))})
        self.db.commit()
        with patch.object(services, 'normalized_adjustments_available', return_value=True) as schema:
            rows = self.rows()
            schema.assert_called_once()
        self.assertEqual({(r['point_code'],r['visit_date']) for r in rows},
                         {('P','2026-09-11'),('OTHER','2026-09-04'),('OTHER','2026-09-11')})

    def test_report_tu_filter_and_status_are_read_only(self):
        self.fixture_3284()
        with (patch.object(services, 'ensure_monthly_submissions_table', side_effect=AssertionError('DDL')),
              patch.object(services, 'ensure_point_adjustments_table', side_effect=AssertionError('DDL'))):
            all_rows = services.get_admin_report_rows(self.db, 2026, 9)
            selected = services.get_admin_report_rows(self.db, 2026, 9, tu='TU-1')
            self.assertEqual(selected, [r for r in all_rows if r['tu']=='TU-1'])
            self.assertEqual(services.get_admin_report_rows(self.db,2026,9,status='submitted'), [])

    def test_export_calculates_sources_once_for_all_sheets(self):
        self.fixture_3284()
        with patch.object(main,'get_intersections_rows',wraps=services.get_intersections_rows) as build:
            response=main.admin_export_overlaps(2026,9,'',main.get_admin_cookie_value(),self.db)
            build.assert_called_once()
        async def read():return [part async for part in response.body_iterator]
        chunks=asyncio.run(read())
        self.assertEqual(len(chunks),1)
        wb=load_workbook(BytesIO(chunks[0]))
        self.assertEqual(wb.sheetnames,['Пересечения','По дням','Итоги'])
        self.assertEqual(wb['Пересечения'].max_row,12)
        self.assertEqual(wb['По дням'].max_row,10)

    def test_reused_styles_match_previous_appearance_and_values(self):
        old,new=Workbook(),Workbook()
        for wb in (old,new):
            for row in [('Header','Long header'),('Сотрудник',123),('',None),('a\nb',-9)]:
                wb.active.append(row)
        ws=old.active
        for cell in ws[1]:
            cell.font=Font(bold=True)
            cell.fill=PatternFill('solid',fgColor='E8F5E9')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for col in ws.columns:
            maximum=0
            for cell in col:
                maximum=max(maximum,len('' if cell.value is None else str(cell.value)))
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            ws.column_dimensions[col[0].column_letter].width=min(max(maximum+2,12),35)
        ws.freeze_panes='A2'
        main.style_sheet(new.active)
        for old_row,new_row in zip(old.active,new.active):
            for a,b in zip(old_row,new_row):
                self.assertEqual(a.value,b.value)
                self.assertEqual(a._style,b._style)
        self.assertEqual(old.active.freeze_panes,new.active.freeze_panes)
        self.assertEqual({k:v.width for k,v in old.active.column_dimensions.items()},
                         {k:v.width for k,v in new.active.column_dimensions.items()})
