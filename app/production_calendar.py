import io
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


CALENDAR_DDL = """
CREATE TABLE IF NOT EXISTS production_calendar (
    calendar_date DATE PRIMARY KEY,
    is_day_off BOOLEAN NOT NULL,
    title TEXT NOT NULL DEFAULT '',
    comment TEXT NOT NULL DEFAULT '',
    year INTEGER NOT NULL,
    source TEXT NOT NULL DEFAULT 'admin',
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
    CHECK (year = EXTRACT(YEAR FROM calendar_date))
)
"""


def ensure_production_calendar_table(db: Session, *, commit: bool = True) -> None:
    db.execute(text(CALENDAR_DDL))
    dialect = getattr(getattr(getattr(db, "bind", None), "dialect", None), "name", "")
    if dialect == "postgresql":
        db.execute(text(
            "ALTER TABLE production_calendar ADD COLUMN IF NOT EXISTS comment TEXT NOT NULL DEFAULT ''"
        ))
    if commit:
        db.commit()


def calendar_day_off(current_date: date, overrides: dict[date, bool]) -> bool:
    """Explicit calendar data overrides the normal Saturday/Sunday rule."""
    if current_date in overrides:
        return bool(overrides[current_date])
    return current_date.weekday() >= 5


def get_calendar_overrides(db: Session, year: int, month: int) -> dict[date, bool]:
    ensure_production_calendar_table(db)
    rows = db.execute(text("""
        SELECT calendar_date, is_day_off FROM production_calendar
        WHERE year=:year AND EXTRACT(MONTH FROM calendar_date)=:month
    """), {"year": year, "month": month}).all()
    return {row[0]: bool(row[1]) for row in rows}


def _bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value or "").strip().lower().replace("ё", "е")
    if normalized in {"1", "true", "yes", "да", "выходной", "нерабочий"}:
        return True
    if normalized in {"0", "false", "no", "нет", "рабочий"}:
        return False
    raise ValueError(f"invalid day-off flag: {value!r}")


def parse_calendar_workbook(file_obj) -> list[dict]:
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = next(iterator)
    except StopIteration:
        raise ValueError("Файл производственного календаря пуст")
    normalized = {
        " ".join(str(value or "").strip().lower().replace("ё", "е").split()): index
        for index, value in enumerate(headers)
    }
    aliases = {
        "calendar_date": ("calendar_date", "дата"),
        "is_day_off": ("is_day_off", "выходной", "выходной день", "нерабочий день"),
        "title": ("title", "название"),
        "source": ("source", "источник"),
        "comment": ("comment", "комментарий", "примечание"),
    }
    indexes = {}
    for canonical, names in aliases.items():
        for name in names:
            if name in normalized:
                indexes[canonical] = normalized[name]
                break
    missing = {"calendar_date", "is_day_off"} - indexes.keys()
    if missing:
        raise ValueError(f"Нет обязательных колонок: {', '.join(sorted(missing))}")
    result, seen = [], set()
    for row_number, row in enumerate(iterator, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        try:
            raw_date = row[indexes["calendar_date"]]
            calendar_date = raw_date.date() if isinstance(raw_date, datetime) else (
                raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)[:10])
            )
            if calendar_date in seen:
                raise ValueError("дата повторяется")
            seen.add(calendar_date)
            result.append({
                "calendar_date": calendar_date,
                "is_day_off": _bool(row[indexes["is_day_off"]]),
                "title": str(row[indexes["title"]] or "").strip() if "title" in indexes else "",
                "comment": str(row[indexes["comment"]] or "").strip() if "comment" in indexes else "",
                "year": calendar_date.year,
                "source": str(row[indexes["source"]] or "admin").strip() if "source" in indexes else "admin",
            })
        except (ValueError, TypeError, IndexError) as exc:
            raise ValueError(f"Ошибка в строке {row_number}: {exc}") from exc
    if not result:
        raise ValueError("В файле нет дат")
    return result


def import_calendar_xlsx(db: Session, file_obj) -> dict:
    rows = parse_calendar_workbook(file_obj)
    try:
        ensure_production_calendar_table(db, commit=False)
        db.execute(text("""
            INSERT INTO production_calendar
                (calendar_date, is_day_off, title, comment, year, source)
            VALUES (:calendar_date, :is_day_off, :title, :comment, :year, :source)
            ON CONFLICT (calendar_date) DO UPDATE SET
                is_day_off=EXCLUDED.is_day_off,
                title=EXCLUDED.title,
                comment=EXCLUDED.comment,
                year=EXCLUDED.year,
                source=EXCLUDED.source,
                updated_at=CURRENT_TIMESTAMP
        """), rows)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"loaded_rows": len(rows), "years": sorted({row["year"] for row in rows})}
