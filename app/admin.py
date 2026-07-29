import hmac
import io
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.security import ADMIN_COOKIE, cookie_secure, make_session, require_admin, safe, verify_csrf


router = APIRouter()
MAX_IMPORT_BYTES = int(os.getenv("MAX_IMPORT_BYTES", str(20 * 1024 * 1024)))


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@router.get("/admin-login", response_class=HTMLResponse)
def admin_login_page():
    return """<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport"
content="width=device-width"><title>Вход администратора</title>
<style>body{font:16px sans-serif;max-width:440px;margin:10vh auto;padding:24px}
input,button{display:block;width:100%;padding:12px;margin:12px 0;box-sizing:border-box}</style>
<h1>Вход администратора</h1><form method="post"><input name="password" type="password"
autocomplete="current-password" required><button>Войти</button></form></html>"""


@router.post("/admin-login")
def admin_login(password: str = Form(...)):
    expected = os.getenv("ADMIN_PASSWORD")
    if not expected or not hmac.compare_digest(password, expected):
        raise HTTPException(status_code=401, detail="Invalid administrator credentials")
    response = RedirectResponse("/admin", status_code=303)
    response.set_cookie(
        ADMIN_COOKIE, make_session("admin", "admin"), httponly=True, secure=cookie_secure(),
        samesite="strict", max_age=12 * 60 * 60,
    )
    return response


REPORT_SQL = """
WITH visit_totals AS (
    SELECT merchant_id, point_code, EXTRACT(YEAR FROM visit_date)::int AS year,
           EXTRACT(MONTH FROM visit_date)::int AS month,
           COUNT(*) FILTER (WHERE slot='DAY') AS day_visits,
           COUNT(*) FILTER (WHERE slot='FULL_INVENT') AS inventories
    FROM visits GROUP BY merchant_id, point_code, year, month
), note_totals AS (
    SELECT merchant_id, point_code, year, month, SUM(amount) AS notes
    FROM point_notes GROUP BY merchant_id, point_code, year, month
), reimbursement_totals AS (
    SELECT merchant_id, point_code, year, month, SUM(amount) AS reimbursements
    FROM point_reimbursements GROUP BY merchant_id, point_code, year, month
), keys AS (
    SELECT merchant_id, point_code, year, month FROM visit_totals
    UNION SELECT merchant_id, point_code, year, month FROM note_totals
    UNION SELECT merchant_id, point_code, year, month FROM reimbursement_totals
)
SELECT m.fio, m.tu, k.point_code, k.year, k.month,
       COALESCE(v.day_visits,0) AS day_visits, COALESCE(v.inventories,0) AS inventories,
       COALESCE(n.notes,0) AS notes, COALESCE(r.reimbursements,0) AS reimbursements,
       EXISTS (SELECT 1 FROM reconciliation_submissions s WHERE s.merchant_id=k.merchant_id
               AND s.year=k.year AND s.month=k.month AND s.reopened_at IS NULL) AS submitted
FROM keys k JOIN merchants m ON m.id=k.merchant_id
LEFT JOIN visit_totals v USING (merchant_id,point_code,year,month)
LEFT JOIN note_totals n USING (merchant_id,point_code,year,month)
LEFT JOIN reimbursement_totals r USING (merchant_id,point_code,year,month)
WHERE (:year IS NULL OR k.year=:year) AND (:month IS NULL OR k.month=:month)
  AND (:tu IS NULL OR m.tu=:tu)
ORDER BY m.fio, k.point_code
"""


def _report_rows(db: Session, year: int | None, month: int | None, tu: str | None):
    return db.execute(text(REPORT_SQL), {"year": year, "month": month, "tu": tu or None}).mappings().all()


@router.get("/admin", response_class=HTMLResponse)
def admin_report(
    request: Request, year: int | None = None, month: int | None = None,
    tu: str | None = None, db: Session = Depends(get_db),
):
    session = require_admin(request)
    rows = _report_rows(db, year, month, tu)
    body = "".join(
        f"<tr><td>{safe(r['fio'])}</td><td>{safe(r['tu'] or '')}</td><td>{safe(r['point_code'])}</td>"
        f"<td>{r['month']:02d}.{r['year']}</td><td>{r['day_visits']}</td><td>{r['inventories']}</td>"
        f"<td>{r['notes']}</td><td>{r['reimbursements']}</td>"
        f"<td>{'Отправлена' if r['submitted'] else 'Черновик'}</td></tr>" for r in rows
    )
    return f"""<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport"
content="width=device-width"><title>Отчёт</title><style>body{{font:14px sans-serif;margin:24px}}
table{{border-collapse:collapse;width:100%}}td,th{{padding:8px;border-bottom:1px solid #ddd;text-align:left}}
form{{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:18px}}input{{padding:8px}}</style>
<h1>Отчёт по сверкам</h1><form><input name="year" type="number" placeholder="Год" value="{year or ''}">
<input name="month" type="number" min="1" max="12" placeholder="Месяц" value="{month or ''}">
<input name="tu" placeholder="ТУ" value="{safe(tu or '')}"><button>Применить</button></form>
<form method="post" action="/admin/import/supplies" enctype="multipart/form-data">
<input type="hidden" name="csrf_token" value="{safe(session["csrf"])}">
<input type="file" name="workbook" accept=".xlsx" required><button>Загрузить поставки</button></form>
<p><a href="/admin/export.xlsx">Выгрузить полный отчёт Excel</a></p>
<table><thead><tr><th>ФИО</th><th>ТУ</th><th>Точка</th><th>Месяц</th><th>Выходы</th>
<th>Инвенты</th><th>Примечания</th><th>Возмещения</th><th>Статус</th></tr></thead>
<tbody>{body}</tbody></table></html>"""


@router.get("/admin/export.xlsx")
def export_admin_report(
    request: Request, year: int | None = None, month: int | None = None,
    tu: str | None = None, db: Session = Depends(get_db),
):
    require_admin(request)
    rows = _report_rows(db, year, month, tu)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Проверка")
    headers = ["ФИО", "ТУ", "Точка", "Год", "Месяц", "Выходы", "Инвенты", "Примечания", "Возмещения", "Статус"]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row["fio"], row["tu"], row["point_code"], row["year"], row["month"],
            row["day_visits"], row["inventories"], float(row["notes"]),
            float(row["reimbursements"]), "Отправлена" if row["submitted"] else "Черновик",
        ])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="reconciliation-report.xlsx"'},
    )


HEADER_ALIASES = {
    "point_code": {"point_code", "точка", "номер точки", "тт"},
    "supply_date": {"supply_date", "дата", "дата поставки"},
    "boxes": {"boxes", "коробки", "количество коробок", "кол-во коробок"},
    "has_supply": {"has_supply", "есть поставка", "поставка"},
}


def _normalized_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("ё", "е").split())


def parse_supply_workbook(data: bytes) -> list[dict]:
    if not data.startswith(b"PK"):
        raise ValueError("The uploaded file is not an XLSX workbook")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        raise ValueError("The workbook is empty")
    indexes: dict[str, int] = {}
    for index, raw in enumerate(raw_headers):
        normalized = _normalized_header(raw)
        for canonical, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                indexes[canonical] = index
    missing = {"point_code", "supply_date", "boxes"} - indexes.keys()
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    parsed: list[dict] = []
    seen: set[tuple[str, date]] = set()
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        try:
            point = "".join(str(row[indexes["point_code"]] or "").split())
            raw_date = row[indexes["supply_date"]]
            supply_date = raw_date.date() if isinstance(raw_date, datetime) else (
                raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)[:10])
            )
            boxes = int(Decimal(str(row[indexes["boxes"]])))
            if not point or boxes < 0:
                raise ValueError
            if "has_supply" in indexes:
                flag = str(row[indexes["has_supply"]] or "").strip().lower()
                if flag in {"false", "0", "нет", "no"}:
                    boxes = 0
            key = (point, supply_date)
            if key in seen:
                raise ValueError("duplicate point/date")
            seen.add(key)
            parsed.append({"point_code": point, "supply_date": supply_date, "boxes": boxes})
        except (ValueError, TypeError, InvalidOperation) as exc:
            raise ValueError(f"Invalid supply data in row {row_number}: {exc}") from exc
        if len(parsed) > 100_000:
            raise ValueError("The workbook contains too many rows")
    if not parsed:
        raise ValueError("The workbook contains no supply rows")
    return parsed


@router.post("/admin/import/supplies")
async def import_supplies(
    request: Request,
    csrf_token: str = Form(...),
    workbook: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    session = require_admin(request)
    verify_csrf(session, csrf_token)
    data = await workbook.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Workbook is too large")
    try:
        rows = parse_supply_workbook(data)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    try:
        db.execute(text("""
            CREATE TEMP TABLE supply_import_stage (
                point_code TEXT NOT NULL, supply_date DATE NOT NULL, boxes INTEGER NOT NULL
            ) ON COMMIT DROP
        """))
        db.execute(text("""
            INSERT INTO supply_import_stage (point_code, supply_date, boxes)
            VALUES (:point_code, :supply_date, :boxes)
        """), rows)
        db.execute(text("""
            UPDATE supplies s SET boxes=stage.boxes
            FROM supply_import_stage stage
            WHERE s.point_code=stage.point_code AND s.supply_date=stage.supply_date
        """))
        db.execute(text("""
            INSERT INTO supplies (point_code, supply_date, boxes)
            SELECT stage.point_code, stage.supply_date, stage.boxes
            FROM supply_import_stage stage
            WHERE NOT EXISTS (
                SELECT 1 FROM supplies s
                WHERE s.point_code=stage.point_code AND s.supply_date=stage.supply_date
            )
        """))
        db.commit()
    except Exception:
        db.rollback()
        raise
    point_count = len({row["point_code"] for row in rows})
    return {"status": "ok", "rows": len(rows), "points": point_count}
